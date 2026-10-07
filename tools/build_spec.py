#!/usr/bin/env python3
"""Compiles a BSP report package into src/spec-<report>.js.

A package is the pair of files the BSP publishes for a report:
  <REPORT>_v<VER>.xsd                  the XML schema of the submission file
  <REPORT>_v<VER>_specification.xlsx   one sheet per schedule plus an "Assertions" sheet

Run:  python tools/build_spec.py "<folder holding the package>" --title "<report name>"
Needs: Python 3.8+, openpyxl.

The specification workbook is slow to open (about two minutes). Pass
--cache FILE to keep the extracted cells between runs while working on this script.
"""
import argparse
import base64
import collections
import glob
import gzip
import json
import os
import pickle
import re
import sys
import xml.etree.ElementTree as ET

XS = '{http://www.w3.org/2001/XMLSchema}'
HERE = os.path.dirname(os.path.abspath(__file__))


# ---- XSD ---------------------------------------------------------------------

def read_xsd(path):
    root = ET.parse(path).getroot()
    types = {}
    for st in root.findall(XS + 'simpleType'):
        res = st.find(XS + 'restriction')
        t = {'base': res.get('base').split(':')[-1]}
        enums = []
        for f in res:
            tag = f.tag[len(XS):]
            if tag == 'enumeration':
                doc = f.find(XS + 'annotation/' + XS + 'documentation')
                enums.append([f.get('value'), (doc.text or '').strip() if doc is not None else ''])
            elif tag == 'pattern':
                t['pattern'] = f.get('value')
            elif tag in ('maxInclusive', 'minInclusive', 'totalDigits', 'fractionDigits', 'minLength', 'maxLength'):
                t[tag] = f.get('value')
        if enums:
            t['enum'] = enums
        types[st.get('name')] = t

    top = root.find(XS + 'element')

    def group(el):
        ct = el.find(XS + 'complexType')
        if ct is None:
            return None, []
        g = ct.find(XS + 'all')
        kind = 'all'
        if g is None:
            g = ct.find(XS + 'sequence')
            kind = 'seq'
        return kind, [c for c in g if c.tag == XS + 'element']

    def leaf(el):
        return [el.get('type'), el.get('minOccurs') == '1']

    def table(el):
        kind, kids = group(el)
        if kind == 'seq':
            item = kids[0]
            _, cells = group(item)
            return {'list': True, 'item': item.get('name'), 'cells': collections.OrderedDict((c.get('name'), leaf(c)) for c in cells)}
        return {'list': False, 'cells': collections.OrderedDict((c.get('name'), leaf(c)) for c in kids)}

    def form_body(kids):
        fields, tables = collections.OrderedDict(), collections.OrderedDict()
        for c in kids:
            if c.get('type'):
                fields[c.get('name')] = leaf(c)
            else:
                tables[c.get('name')] = table(c)
                tables[c.get('name')]['req'] = c.get('minOccurs') == '1'
        return fields, tables

    header, forms = None, collections.OrderedDict()
    _, kids = group(top)
    for el in kids:
        name = el.get('name')
        kind, sub = group(el)
        if name == 'Header':
            header = [[c.get('name'), c.get('type')] for c in sub]
            continue
        f = {'req': el.get('minOccurs') == '1'}
        if kind == 'seq':
            item = sub[0]
            _, inner = group(item)
            f['item'] = item.get('name')
            f['fields'], f['tables'] = form_body(inner)
        else:
            f['fields'], f['tables'] = form_body(sub)
        forms[name] = f

    unique = []
    for u in top.findall(XS + 'unique'):
        sel = u.find(XS + 'selector').get('xpath').replace('ns:', '').split('/')
        unique.append([sel[0], u.find(XS + 'field').get('xpath').replace('ns:', '')])
    return {'ns': root.get('targetNamespace'), 'root': top.get('name'), 'types': types,
            'header': header, 'forms': forms, 'unique': unique}


# ---- specification workbook --------------------------------------------------

def read_workbook(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=False)
    raw = {'sheets': collections.OrderedDict(), 'assertions': []}
    for ws in wb.worksheets:
        if ws.title == 'Assertions':
            raw['assertions'] = [list(r) for r in ws.iter_rows(values_only=True)]
            continue
        grid, notes = [], {}
        for row in ws.iter_rows():
            grid.append([c.value for c in row])
            for c in row:
                if c.comment:
                    notes[(c.row - 1, c.column - 1)] = c.comment.text
        merged = [(m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1) for m in ws.merged_cells.ranges]
        raw['sheets'][ws.title] = {'grid': grid, 'notes': notes, 'merged': merged}
    return raw


def clean(s):
    return re.sub(r'\s+', ' ', str(s)).strip() if s is not None else ''


CELL = re.compile(r'^([A-Za-z][A-Za-z0-9_]*)\n(.+)$', re.S)
PCT = re.compile(r'^- <([0-9.]+)%> -$')


def note_parts(text):
    """Returns (formula, condition) from a cell note."""
    formula = cond = None
    for part in re.split(r'\n\s*\n', text or ''):
        part = part.strip()
        m = re.match(r'^Condition formula:\s*(.+)$', part, re.S)
        if m:
            cond = clean(m.group(1))
            continue
        m = re.match(r'^Formula:\s*(.+)$', part, re.S)
        if m:
            formula = clean(m.group(1))
    return formula, cond


def read_form(sheet):
    grid, notes = sheet['grid'], sheet['notes']
    span = {}
    for r0, c0, r1, c1 in sheet['merged']:
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                span[(r, c)] = (r0, c0)

    def at(r, c):
        r, c = span.get((r, c), (r, c))
        return grid[r][c] if r < len(grid) and c < len(grid[r]) else None

    def a(r):
        return clean(grid[r][0]) if grid[r] else ''

    def c3(r):
        return grid[r][2] if len(grid[r]) > 2 else None

    form = {'node': clean(grid[0][2]).split('/')[0], 'title': clean(grid[1][2]), 'z': None, 'fields': [], 'tables': []}
    for r in range(len(grid)):
        if a(r) == 'Z Axis Field' and c3(r) and c3(r) != '- none -':
            form['z'] = {'field': str(c3(r)).split('\n')[0], 'title': clean(c3(r + 1))}
        if a(r) == 'Field XML node':
            form['fields'].append({'name': str(c3(r)).split('\n')[0], 'title': clean(c3(r + 1))})

    starts = [r for r in range(len(grid)) if a(r) == 'Table XML Node']
    ends = starts[1:] + [min([r for r in range(len(grid)) if a(r) == 'Field XML node'] + [len(grid)])]
    for start, end in zip(starts, ends):
        name, kind = str(c3(start)).split('\n')
        kind = {'(SIMPLE)': 'S', '(Y-AXIS)': 'Y', '(X-AXIS)': 'X'}[kind.strip()]
        title = clean(c3(start + 1))
        t = {'name': name, 'kind': kind, 'title': '' if title == '- none -' else title,
             'rows': [], 'cols': [], 'cells': collections.OrderedDict()}

        # the row of column codes (C0010, C0020, ...) and the header rows above it
        code_row = None
        if kind != 'X':
            for r in range(start + 2, end):
                if any(isinstance(v, str) and re.fullmatch(r'C\d{4}', v) for v in grid[r]):
                    code_row = r
                    break
        col_labels = {}
        if code_row is not None:
            for c, v in enumerate(grid[code_row]):
                if isinstance(v, str) and re.fullmatch(r'C\d{4}', v):
                    parts = []
                    for r in range(start + 2, code_row):
                        h = clean(at(r, c))
                        if h and not re.fullmatch(r'\d{15,}', h) and h not in parts:
                            parts.append(h)
                    col_labels[v] = ' / '.join(parts)

        row_seen = set()
        for r in range((code_row if code_row is not None else start + 1) + 1, end):
            row = grid[r]
            # the row code sits in its own cell, left of the data cells
            code_at = next((c for c, v in enumerate(row[:6]) if isinstance(v, str) and re.fullmatch(r'R\d{4}|[A-Z]{2,}[A-Z0-9_]*', v)
                            and c > 0 and (kind != 'Y')), None)
            if kind == 'S' and code_at is not None and not re.fullmatch(r'R\d{4}', row[code_at]):
                code_at = None
            if code_at is not None and row[code_at] not in row_seen:
                left = [clean(v) for v in row[:code_at]]
                acct = next((v for v in left if re.fullmatch(r'\d{3,}', v)), '')
                label = ' '.join(v for v in left if v and not re.fullmatch(r'\d{3,}', v))
                t['rows'].append([row[code_at], label, acct])
                row_seen.add(row[code_at])
            for c, v in enumerate(row):
                if not isinstance(v, str):
                    continue
                m = CELL.match(v)
                if not m:
                    continue
                code, tag = m.group(1), m.group(2).strip()
                formula, cond = note_parts(notes.get((r, c)))
                cell = {'tag': tag, 'cond': cond}
                pct = PCT.match(tag)
                if pct:
                    cell['const'] = float(pct.group(1)) / 100
                elif tag == '- autogenerated -':
                    cell['formula'] = formula
                t['cells'][code] = cell

        if kind == 'S':
            used = set(code[5:] for code in t['cells'])
            t['cols'] = [[c, col_labels[c]] for c in col_labels if c in used]
        elif kind == 'Y':
            t['cols'] = [[c, col_labels.get(c, '')] for c in t['cells']]
            t['rows'] = []
        else:
            t['cols'] = []
        form['tables'].append(t)
    return form


# ---- merge -------------------------------------------------------------------

def build(xsd, raw):
    stats = collections.Counter()
    type_ids, types_out = {}, []
    for name, t in xsd['types'].items():
        type_ids[name] = len(types_out)
        o = {'b': {'decimal': 'd', 'integer': 'i', 'string': 's', 'date': 'D'}[t['base']], 'n': re.sub(r'(___|_)\d+$', '', name)}
        for k, short in (('minInclusive', 'min'), ('maxInclusive', 'max'), ('totalDigits', 'td'), ('fractionDigits', 'fd'),
                         ('minLength', 'minLen'), ('maxLength', 'maxLen'), ('pattern', 'pat'), ('enum', 'enum')):
            if k in t:
                o[short] = t[k]
        types_out.append(o)

    conds, cond_ids = [], {}

    def cond_id(text):
        if text is None:
            return None
        if text not in cond_ids:
            cond_ids[text] = len(conds)
            conds.append(text)
        return cond_ids[text]

    def cell_out(spec_cell, xsd_leaf, where):
        """[kind, value, condition]: kind 0 = reported (value: type id, required),
        1 = calculated (value: formula), 2 = fixed number."""
        cond = cond_id(spec_cell['cond']) if spec_cell else None
        if xsd_leaf is not None:
            if spec_cell is None:
                stats['in XSD, not in specification'] += 1
            elif 'formula' in spec_cell or 'const' in spec_cell:
                stats['in XSD but calculated in specification'] += 1
            out = [0, type_ids[xsd_leaf[0]] if not xsd_leaf[1] else [type_ids[xsd_leaf[0]]]]
        elif 'const' in spec_cell:
            out = [2, spec_cell['const']]
        elif spec_cell.get('formula'):
            out = [1, spec_cell['formula']]
        elif spec_cell['tag'] == '- autogenerated -':
            stats['calculated cell without a formula'] += 1
            return None
        else:
            stats['reported in specification, not in XSD'] += 1
            return None
        if cond is not None:
            out.append(cond)
        return out

    forms_out = []
    spec_forms = collections.OrderedDict()
    for title, sheet in raw['sheets'].items():
        f = read_form(sheet)
        spec_forms[f['node']] = f
    order = list(xsd['forms'])
    for name in spec_forms:
        if name not in xsd['forms']:
            # a schedule the BSP derives from others; keep it next to its sheet neighbours
            names = list(spec_forms)
            prev = next((n for n in reversed(names[:names.index(name)]) if n in order), None)
            order.insert(order.index(prev) + 1 if prev else 0, name)
    for name in order:
        sf, xf = spec_forms.get(name), xsd['forms'].get(name)
        if sf is None:
            stats['form in XSD, not in specification'] += 1
            sf = {'title': '', 'z': None, 'fields': [], 'tables': []}
        fo = {'n': name, 't': sf['title']}
        if xf is not None:
            fo['x'] = 1
            if xf['req']:
                fo['r'] = 1
            if xf.get('item'):
                zf = next(iter(xf['fields']))
                fo['z'] = {'item': xf['item'], 'f': zf, 'ty': type_ids[xf['fields'][zf][0]], 't': (sf['z'] or {}).get('title', '')}
            else:
                titles = {x['name']: x['title'] for x in sf['fields']}
                if xf['fields']:
                    fo['f'] = [[n, type_ids[l[0]], titles.get(n, '')] for n, l in xf['fields'].items()]
        tables = []
        spec_tables = collections.OrderedDict((t['name'], t) for t in sf['tables'])
        xsd_tables = xf['tables'] if xf else {}
        for tn in list(spec_tables) + [n for n in xsd_tables if n not in spec_tables]:
            st, xt = spec_tables.get(tn), xsd_tables.get(tn)
            if st is None:
                stats['table in XSD, not in specification'] += 1
                st = {'kind': 'Y' if xt['list'] else 'S', 'title': '', 'rows': [], 'cols': [], 'cells': {}}
            to = {'n': tn, 'k': st['kind'], 't': st['title']}
            if xt is not None:
                to['x'] = 1
                if xt.get('req'):
                    to['r'] = 1
                if xt['list']:
                    to['item'] = xt['item']
            xcells = xt['cells'] if xt else {}
            cells = collections.OrderedDict()
            for code in list(st['cells']) + [c for c in xcells if c not in st['cells']]:
                out = cell_out(st['cells'].get(code), xcells.get(code), (name, tn, code))
                if out is not None:
                    cells[code] = out
            to['c'] = cells
            if st['kind'] == 'S':
                to['rows'] = st['rows']
                to['cols'] = st['cols']
            elif st['kind'] == 'Y':
                to['cols'] = st['cols']
            else:
                to['rows'] = st['rows']
            tables.append(to)
            stats['tables'] += 1
            stats['cells'] += len(cells)
        fo['tb'] = tables
        forms_out.append(fo)

    rules, seen_rules = [], set()
    head = raw['assertions'][0]
    assert head[:9] == ['Code', 'Message', 'Left formula', 'Operator', 'Right formula', 'Tolerance', 'Loop field', 'Precondition', 'Severity'], head
    for r in raw['assertions'][1:]:
        if r[0] is None:
            continue
        code, msg = clean(r[0]), clean(r[1])
        # some messages repeat their own code in front; the code is shown anyway
        msg = re.sub(r'^STG1-[A-Za-z0-9_.-]+: ', '', msg) if code.startswith('STG1-') else msg
        if msg.startswith(code + ': '):
            msg = msg[len(code) + 2:]
        elif msg.startswith(code + ' '):
            msg = msg[len(code) + 1:]
        rule = [code, msg, clean(r[2]), r[3], clean(r[4]), r[5] or 0,
                clean(r[6]) or 0, clean(r[7]) or 0, 1 if r[8] == 'Warning' else 0]
        key = json.dumps(rule)
        if key in seen_rules:
            stats['duplicate rules dropped'] += 1
            continue
        seen_rules.add(key)
        rules.append(rule)
    stats['rules'] = len(rules)
    stats['forms'] = len(forms_out)

    return {'ns': xsd['ns'], 'root': xsd['root'], 'types': types_out,
            'header': [[n, type_ids[t]] for n, t in xsd['header']],
            'unique': xsd['unique'], 'conds': conds, 'forms': forms_out, 'rules': rules}, stats


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('folder', help='folder holding the .xsd and the _specification.xlsx')
    ap.add_argument('--title', default='', help='name of the report as the BSP writes it, shown on the page')
    ap.add_argument('--cache', help='pickle file for the extracted workbook cells')
    ap.add_argument('--out', help='output file (default: src/spec-<report>.js)')
    args = ap.parse_args()

    xsd_path = glob.glob(os.path.join(args.folder, '*.xsd'))
    spec_path = glob.glob(os.path.join(args.folder, '*_specification.xlsx'))
    if len(xsd_path) != 1 or len(spec_path) != 1:
        sys.exit('Expected exactly one .xsd and one *_specification.xlsx in ' + args.folder)

    if args.cache and os.path.exists(args.cache):
        raw = pickle.load(open(args.cache, 'rb'))
    else:
        print('reading', os.path.basename(spec_path[0]), '...')
        raw = read_workbook(spec_path[0])
        if args.cache:
            pickle.dump(raw, open(args.cache, 'wb'))
    xsd = read_xsd(xsd_path[0])
    spec, stats = build(xsd, raw)
    m = re.match(r'(.+)_v([0-9.]+)\.xsd$', os.path.basename(xsd_path[0]))
    spec['report'], spec['version'] = m.group(1), m.group(2)
    spec['title'] = args.title
    out = args.out or os.path.join(HERE, '..', 'src', 'spec-%s.js' % spec['report'].lower())

    # What the page needs to recognise a file before it unpacks the definition.
    meta = collections.OrderedDict((k, spec[k]) for k in ('report', 'version', 'title', 'root', 'ns'))
    meta['sheets'] = [f['n'] for f in spec['forms'] if f.get('x')]
    text = json.dumps(spec, ensure_ascii=True, separators=(',', ':'))
    meta['blob'] = base64.b64encode(gzip.compress(text.encode('utf-8'), 9, mtime=0)).decode('ascii')
    lines = [
        '// Generated by tools/build_spec.py from %s and %s. Do not edit.' % (
            os.path.basename(xsd_path[0]), os.path.basename(spec_path[0])),
        '// "blob" is the gzip + base64 of the compiled report definition (%d KB of JSON).' % (len(text) // 1024),
        '(function (root) {',
        "  'use strict';",
        '  var BSPV = root.BSPV = root.BSPV || {};',
        '  (BSPV.specs = BSPV.specs || {})[%s] = %s;' % (
            json.dumps(spec['report']), json.dumps(meta, ensure_ascii=True, separators=(',', ':'))),
        "})(typeof self !== 'undefined' ? self : globalThis);",
        ''
    ]
    with open(out, 'wb') as f:
        f.write(chr(10).join(lines).encode('ascii'))
    for k in sorted(stats):
        print('%7d  %s' % (stats[k], k))
    print('wrote %s: %d KB (%d KB before compression)' % (os.path.relpath(out), len(meta['blob']) // 1024, len(text) // 1024))
    return spec


if __name__ == '__main__':
    main()
